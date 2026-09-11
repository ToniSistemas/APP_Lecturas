import base64
import csv
import io
import unicodedata

from odoo import _, fields, models
from odoo.exceptions import UserError, ValidationError


class WaterMeterImport(models.TransientModel):
    _name = 'water.meter.import'
    _description = 'Importar contadores de agua'

    file = fields.Binary(string='Archivo', required=True)
    filename = fields.Char(string='Nombre del archivo')
    period_id = fields.Many2one('water.period', string='Período', required=True)
    import_readings = fields.Boolean(string='Importar lecturas', default=False, readonly=True)
    confirmation_required = fields.Boolean(readonly=True)
    warning_message = fields.Text(string='Lecturas diferentes', readonly=True)

    _column_fields = {
        'Ruta': 'name',
        'Contador': 'meter_number',
        'Nombre': 'owner_name',
        'Abonado': 'subscriber',
        'Referencia catastral': 'cadastral_ref',
        'Dirección': 'address',
        'Calle': 'street',
        'Ubicación': 'street_number',
        'C.P.': 'zip',
        'Municipio': 'municipality',
        'Tipo de contador': 'meter_type',
        'Lectura anterior': 'reading_previous',
        'Lectura actual': 'reading_current',
        'Consumo': 'consumption',
    }
    _required_headers = {'Ruta', 'Contador', 'Nombre', 'Abonado'}

    _meter_types = {
        'DOM': 'dom',
        'DOMESTICO': 'dom',
        'DOMÉSTICO': 'dom',
        'ASIM': 'asim',
        'ASIMILADO': 'asim',
        'NDOM': 'ndom',
        'NO DOMESTICO': 'ndom',
        'NO DOMÉSTICO': 'ndom',
        'ESP': 'esp',
        'ESPECIAL': 'esp',
    }

    @staticmethod
    def _normalize_header(value):
        text = ''.join(str(value or '').replace('\xa0', ' ').split()).casefold()
        return ''.join(
            character for character in unicodedata.normalize('NFKD', text)
            if not unicodedata.combining(character)
        )

    def _read_rows(self):
        self.ensure_one()
        content = base64.b64decode(self.file or b'')
        filename = (self.filename or '').lower()
        if filename.endswith('.csv'):
            text = content.decode('utf-8-sig')
            return list(csv.DictReader(io.StringIO(text)))
        if filename.endswith('.xlsx'):
            try:
                from openpyxl import load_workbook
            except ImportError as error:
                raise UserError(_("Para importar archivos XLSX debe instalarse la librería Python 'openpyxl'.")) from error
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                return []
            headers = [self._normalize_header(header) for header in headers]
            return [
                dict(zip(headers, values))
                for values in rows
                if any(value not in (None, '') for value in values)
            ]
        raise UserError(_('El archivo debe tener formato CSV o XLSX.'))

    def _confirmation_action(self):
        return {
            'type': 'ir.actions.act_window',
            'name': _('Confirmar importación del censo'),
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_import_confirmed(self):
        self.ensure_one()
        return self.with_context(confirm_reading_mismatches=True).action_import()

    def action_import(self):
        self.ensure_one()
        if self.period_id.state == 'closed':
            raise UserError(_('No se puede importar el censo en un período cerrado.'))
        rows = self._read_rows()
        if not rows:
            raise UserError(_('El archivo no contiene contadores para importar.'))

        normalized_rows = [
            {self._normalize_header(header): value for header, value in row.items()}
            for row in rows
        ]
        required_headers = {self._normalize_header(header) for header in self._required_headers}
        missing_headers = required_headers - set(normalized_rows[0])
        if missing_headers:
            raise ValidationError(
                _('Faltan columnas obligatorias: %s') % ', '.join(sorted(missing_headers)).title()
            )

        available_fields = {
            field
            for header, field in self._column_fields.items()
            if self._normalize_header(header) in normalized_rows[0]
        }
        imported_rows = []
        meter_numbers = set()
        subscribers = set()
        routes = set()
        for row_number, row in enumerate(normalized_rows, start=2):
            values = {
                field: row.get(self._normalize_header(header), '')
                for header, field in self._column_fields.items()
                if field in available_fields
            }
            values = {
                field: str(value).strip() if value is not None else ''
                for field, value in values.items()
            }
            if (
                not values['name'] or not values['meter_number']
                or not values['owner_name'] or not values['subscriber']
            ):
                raise ValidationError(
                    _('Fila %s: Ruta, Contador, Nombre y Abonado son obligatorios.') % row_number
                )
            if values['name'] in routes:
                raise ValidationError(_('Fila %s: la ruta %s está repetida en el archivo.') % (
                    row_number, values['name']))
            if values['meter_number'] in meter_numbers:
                raise ValidationError(_('Fila %s: el contador %s está repetido en el archivo.') % (
                    row_number, values['meter_number']))
            if values['subscriber'] in subscribers:
                raise ValidationError(_('Fila %s: el abonado %s está repetido en el archivo.') % (
                    row_number, values['subscriber']))
            if 'meter_type' in values:
                meter_type = values['meter_type'].upper()
                if meter_type and meter_type not in self._meter_types:
                    raise ValidationError(
                        _('Fila %s: Tipo de contador debe ser DOM, ASIM, NDOM o ESP.') % row_number
                    )
                values['meter_type'] = self._meter_types.get(meter_type, False)
            previous_reading = values.pop('reading_previous', False)
            if previous_reading:
                try:
                    previous_reading = int(previous_reading)
                except (TypeError, ValueError) as error:
                    raise ValidationError(_('Fila %s: Lectura anterior debe ser un número entero.') % row_number) from error
                if previous_reading < 0:
                    raise ValidationError(_('Fila %s: Lectura anterior no puede ser negativa.') % row_number)
            else:
                previous_reading = False
            current_reading = values.pop('reading_current', False)
            imported_consumption = values.pop('consumption', False)
            if current_reading:
                try:
                    current_reading = int(current_reading)
                except (TypeError, ValueError) as error:
                    raise ValidationError(_('Fila %s: Lectura actual debe ser un número entero.') % row_number) from error
                if current_reading < 0:
                    raise ValidationError(_('Fila %s: Lectura actual no puede ser negativa.') % row_number)
            else:
                current_reading = False
            if imported_consumption:
                try:
                    imported_consumption = int(imported_consumption)
                except (TypeError, ValueError) as error:
                    raise ValidationError(_('Fila %s: Consumo debe ser un número entero.') % row_number) from error
            else:
                imported_consumption = False
            routes.add(values['name'])
            meter_numbers.add(values['meter_number'])
            subscribers.add(values['subscriber'])
            imported_rows.append((values, previous_reading, current_reading, imported_consumption))

        Meter = self.env['water.meter'].with_context(active_test=False)
        existing_meters = Meter.search([('subscriber', 'in', list(subscribers))])
        duplicate_subscribers = {
            subscriber
            for subscriber in subscribers
            if len(existing_meters.filtered(lambda meter: meter.subscriber == subscriber)) > 1
        }
        if duplicate_subscribers:
            raise ValidationError(
                _('Hay varios contadores guardados para estos abonados: %s')
                % ', '.join(sorted(duplicate_subscribers))
            )
        existing_by_subscriber = {meter.subscriber: meter for meter in existing_meters}
        existing_by_route = {
            meter.name: meter
            for meter in Meter.search([('name', 'in', list(routes))])
        }
        existing_by_number = {
            meter.meter_number: meter
            for meter in Meter.search([('meter_number', 'in', list(meter_numbers))])
        }
        for values, _previous_reading, _current_reading, _consumption in imported_rows:
            subscriber_meter = existing_by_subscriber.get(values['subscriber'])
            number_owner = existing_by_number.get(values['meter_number'])
            if number_owner and number_owner != subscriber_meter:
                raise ValidationError(
                    _('El contador %(meter)s ya pertenece al abonado %(subscriber)s.') % {
                        'meter': values['meter_number'],
                        'subscriber': number_owner.subscriber or '-',
                    }
                )

        previous_period = self.period_id._get_previous_period()
        previous_by_meter = {}
        if previous_period:
            previous_by_meter = {
                reading.meter_id.id: reading.reading_current
                for reading in self.env['water.reading'].search([
                    ('period_id', '=', previous_period.id),
                    ('meter_id', 'in', existing_meters.ids),
                ])
            }
        mismatches = []
        for values, imported_previous, _current_reading, _consumption in imported_rows:
            meter = existing_by_subscriber.get(values['subscriber'])
            if not meter or imported_previous is False or meter.id not in previous_by_meter:
                continue
            stored_current = previous_by_meter[meter.id]
            if imported_previous != stored_current:
                mismatches.append(
                    _('%(meter)s (%(subscriber)s): guardada %(stored)s, Excel %(imported)s') % {
                        'meter': meter.meter_number,
                        'subscriber': meter.subscriber or meter.owner_name or '-',
                        'stored': stored_current,
                        'imported': imported_previous,
                    }
                )
        if mismatches and not self.env.context.get('confirm_reading_mismatches'):
            shown_mismatches = mismatches[:50]
            if len(mismatches) > 50:
                shown_mismatches.append(_('... y %s diferencias más.') % (len(mismatches) - 50))
            self.write({
                'confirmation_required': True,
                'warning_message': _(
                    'La lectura anterior del Excel no coincide con la lectura actual de '
                    '%(period)s para estos contadores:\n\n%(differences)s'
                ) % {
                    'period': previous_period.name,
                    'differences': '\n'.join(shown_mismatches),
                },
            })
            return self._confirmation_action()

        meters = []
        created_count = 0
        updated_count = 0
        for values, previous_reading, current_reading, imported_consumption in imported_rows:
            meter = existing_by_subscriber.get(values['subscriber'])
            route_owner = existing_by_route.get(values['name'])
            if route_owner and route_owner != meter:
                raise ValidationError(
                    _('La ruta %(route)s ya pertenece al contador %(meter)s.') % {
                        'route': values['name'],
                        'meter': route_owner.meter_number,
                    }
                )
            if meter:
                previous_meter_number = meter.meter_number
                number_changed = values['meter_number'] != previous_meter_number
                update_values = dict(values, active=True)
                changed_values = {
                    field: value
                    for field, value in update_values.items()
                    if meter[field] != value
                }
                if changed_values:
                    meter.write(changed_values)
                    updated_count += 1
            else:
                meter = Meter.with_context(skip_initial_reading=True).create(dict(values, active=True))
                previous_meter_number = False
                number_changed = False
                existing_by_subscriber[meter.subscriber] = meter
                existing_by_route[meter.name] = meter
                created_count += 1
            meter_event = {
                'new_meter': number_changed,
                'previous_meter_number': previous_meter_number if number_changed else False,
                'new_meter_number': meter.meter_number if number_changed else False,
                'meter_number': meter.meter_number,
            }
            meters.append((meter, previous_reading, meter_event))

        Reading = self.env['water.reading']
        readings_by_meter = {
            reading.meter_id.id: reading
            for reading in Reading.search([
                ('period_id', '=', self.period_id.id),
                ('meter_id', 'in', [meter.id for meter, _previous, _event in meters]),
            ])
        }
        reading_count = 0
        for meter, previous_reading, meter_event in meters:
            reading = readings_by_meter.get(meter.id)
            if reading:
                reading_values = {}
                if previous_reading is not False and reading.reading_previous != previous_reading:
                    reading_values['reading_previous'] = previous_reading
                if self.import_readings and current_reading is not False:
                    reading_values['reading_current'] = current_reading
                if meter_event['new_meter']:
                    reading_values.update(meter_event)
                if reading_values:
                    reading.with_context(skip_meter_replacement=True).write(reading_values)
                if self.import_readings and imported_consumption is not False:
                    if reading.difference != imported_consumption:
                        raise ValidationError(
                            _('El consumo del contador %s no coincide con el Excel.') % meter.meter_number
                        )
                continue
            reading_values = {
                'meter_id': meter.id,
                'period_id': self.period_id.id,
                'date': fields.Date.context_today(self),
                **meter_event,
            }
            if previous_reading is not False:
                reading_values['reading_previous'] = previous_reading
            if self.import_readings and current_reading is not False:
                reading_values['reading_current'] = current_reading
            reading = Reading.create(reading_values)
            if self.import_readings and imported_consumption is not False:
                if reading.difference != imported_consumption:
                    raise ValidationError(
                        _('El consumo del contador %s no coincide con el Excel.') % meter.meter_number
                    )
            reading_count += 1

        if self.period_id.state == 'draft':
            self.period_id.state = 'open'

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Importación completada'),
                'message': _(
                    '%(created)s contadores creados, %(updated)s actualizados y '
                    '%(readings)s añadidos al período %(period)s.'
                ) % {
                    'created': created_count,
                    'updated': updated_count,
                    'readings': reading_count,
                    'period': self.period_id.name,
                },
                'type': 'success',
                'sticky': False,
            },
        }